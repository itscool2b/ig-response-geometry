"""Export the canonical manuscript to the journal's Online Word representation.

Pandoc supplies structural conversion and native tables. The Online citation
strings come from the actual compiled bibliography. Equations intentionally
remain TeX source as requested by the journal's LaTeX submission instructions.
No network request is made by this exporter.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import subprocess
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def digest(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def strip_comments(text):
    # Full-line review markers must not become blank paragraphs inside captions.
    text = re.sub(r"(?m)^[ \t]*%[^\n]*(?:\n|$)", "", text)
    return re.sub(r"(?<!\\)%[^\n]*", "", text)


def strip_pdf_layout_definitions(text):
    """Keep Pandoc's structural headings instead of expanding PDF-only macros."""
    def remove(match):
        lines = [line.strip() for line in match.group(1).splitlines() if line.strip()]
        if not lines or not all(re.match(r"\\renewcommand\\(?:section|subsection|@maketitle)\{", line) for line in lines):
            raise ValueError("Unreviewed makeatletter block in Online export")
        return ""
    return re.sub(r"\\makeatletter\b(.*?)\\makeatother\b", remove, text, flags=re.S)


def select_identified(text):
    """Evaluate only the manuscript's reviewed anonymity conditionals."""
    text = re.sub(r"\\newif\\ifanonymous\s*", "", text)
    token = re.compile(r"\\ifdefined\\TMLRAnonymous\b|\\ifanonymous\b|\\else\b|\\fi\b")
    stack, active, out, start = [], True, [], 0
    for match in token.finditer(text):
        if active:
            out.append(text[start:match.start()])
        word = match.group()
        if word.startswith(r"\if"):
            stack.append((active, False))
            active = False
        elif word == r"\else":
            if not stack or stack[-1][1]:
                raise ValueError("Unreviewed manuscript conditional")
            parent, _ = stack[-1]
            stack[-1] = (parent, True)
            active = parent
        else:
            if not stack:
                raise ValueError("Unmatched manuscript conditional")
            active = stack.pop()[0]
        start = match.end()
    if stack:
        raise ValueError("Unclosed manuscript conditional")
    if active:
        out.append(text[start:])
    return re.sub(r"\\anonymous(?:true|false)\b", "", "".join(out))


def expand_inputs(text, paper: Path, seen=None):
    seen = set() if seen is None else set(seen)
    def replace(match):
        name = match.group(1)
        path = (paper / name).with_suffix(".tex") if not Path(name).suffix else paper / name
        path = path.resolve()
        if not path.is_relative_to(paper.resolve()) or path in seen:
            raise ValueError("Unsafe or cyclic manuscript input")
        return expand_inputs(strip_comments(path.read_text(encoding="utf-8")), paper, seen | {path})
    return re.sub(r"\\input\{([^}]+)\}", replace, text)


def bibliography(bbl_text):
    matches = list(re.finditer(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", bbl_text))
    if not matches:
        raise ValueError("No bibliography entries in compiled bbl")
    result = {}
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else bbl_text.index(r"\end{thebibliography}", match.end())
        value = bbl_text[match.end():end].strip()
        result[match.group(1)] = value
    return result


def normalize_table_stacks(text):
    """Keep multiline header text within its cell before Pandoc parses rows."""
    # Native table formatting is applied below. Pandoc's column-declaration
    # parser can consume a leading numeral after this PDF-only modifier.
    text = text.replace(r">{\raggedright\arraybackslash}", "")
    text = re.sub(r"\\shortstack(?:\[[lcr]\])?\{([^{}]*)\}",
                  lambda m: re.sub(r"\\\\\s*", " ", m.group(1)), text)
    if r"\shortstack" in text:
        raise ValueError("Unreviewed nested shortstack in Online export")
    return text


def number_headings(blocks, labels):
    """Retain source section numbers used by cross-references in the text."""
    section, subsection, current = 0, 0, None
    for block in blocks:
        if block.get("t") != "Header":
            continue
        level, attributes, inlines = block["c"]
        if "unnumbered" in attributes[1]:
            continue
        if level == 1:
            subsection = 0
            number = labels.get(attributes[0], "")
            if re.fullmatch(r"[A-Z]", number):
                current = number
                prefix = "Appendix " + number + ". "
            else:
                section += 1
                current = str(section)
                prefix = current + ". "
        elif level == 2 and current is not None:
            subsection += 1
            prefix = current + "." + str(subsection) + " "
        else:
            continue
        block["c"][2] = text_inlines(prefix) + inlines


def pandoc_run(executable, text, source_format, target_format, *extra):
    result = subprocess.run([str(executable), "--from", source_format, "--to", target_format, *map(str, extra)],
                            input=text, text=True, encoding="utf-8", capture_output=True, check=True)
    return result.stdout, result.stderr


def inline_text(inlines):
    out = []
    for node in inlines:
        t, c = node.get("t"), node.get("c")
        if t == "Str":
            out.append(c)
        elif t in {"Space", "SoftBreak", "LineBreak"}:
            out.append(" ")
        elif t in {"Emph", "Strong", "Superscript", "Subscript", "SmallCaps"}:
            out.append(inline_text(c))
        elif t in {"Link", "Span"}:
            out.append(inline_text(c[1]))
        elif t == "Math":
            out.append(c[1])
        else:
            raise ValueError("Unreviewed plain-text inline " + str(t))
    return "".join(out)


def text_inlines(text):
    chunks = re.split(r"(\s+)", text)
    return [{"t": "Space"} if chunk.isspace() else {"t": "Str", "c": chunk} for chunk in chunks if chunk]


def prepare_blocks(blocks):
    """Give display math its own paragraph before converting it to literal TeX."""
    result = []
    for block in blocks:
        if block.get("t") in {"Para", "Plain"}:
            pending = []
            for item in block["c"]:
                if item.get("t") == "Math" and item["c"][0]["t"] == "DisplayMath":
                    if pending:
                        result.append({"t": block["t"], "c": pending})
                    result.append({"t": "Para", "c": [item]})
                    pending = []
                else:
                    pending.append(item)
            if pending:
                result.append({"t": block["t"], "c": pending})
        else:
            result.append(block)
    return result


def transform_ast(value, references, metrics):
    if isinstance(value, list):
        result = []
        for item in value:
            replacement = transform_ast(item, references, metrics)
            if isinstance(item, dict) and item.get("t") in {"Cite", "Math", "RawBlock", "RawInline"} and isinstance(replacement, list):
                result.extend(replacement)
            else:
                result.append(replacement)
        return result
    if not isinstance(value, dict):
        return value
    kind, content = value.get("t"), value.get("c")
    if kind == "Cite":
        citations, _ = content
        out = [{"t": "Space"}]
        for index, citation in enumerate(citations):
            key = citation["citationId"]
            if key not in references:
                raise ValueError("Citation absent from compiled bibliography " + key)
            if index:
                out.extend([{"t": "Superscript", "c": [{"t": "Str", "c": ","}]}, {"t": "Space"}])
            out += [{"t": "Str", "c": "(("}] + copy.deepcopy(references[key]) + [{"t": "Str", "c": "))"}]
            metrics["citation_uses"].append(key)
        return out
    if kind == "Math":
        display = content[0]["t"] == "DisplayMath"
        metrics["display_equations" if display else "inline_equations"] += 1
        delimiter = ("" if content[1].lstrip().startswith(r"\begin{") else "$$") if display else "$"
        # Literal TeX is intentional for this specific venue export.
        return text_inlines(delimiter + content[1] + delimiter)
    if kind == "Image":
        image = copy.deepcopy(value)
        target = image["c"][2][0]
        if target.endswith(".pdf"):
            image["c"][2][0] = target[:-4] + ".png"
        metrics["images"].append(image["c"][2][0])
        return image
    if kind == "Table":
        metrics["tables"] += 1
    if kind in {"RawBlock", "RawInline"} and content[0] == "latex":
        allowed = (r"\centering", r"\clearpage", r"\FloatBarrier", r"\appendix", r"\newpage", r"\hfill", r"\noindent", r"\medskip", r"\smallskip", r"\vspace", r"\setlength", r"\renewcommand", r"\normalsize", r"\small", r"\footnotesize")
        if content[1].strip().startswith(allowed):
            return []
        raise ValueError("Unconverted meaningful TeX in Online export " + content[1][:180])
    return {key: transform_ast(val, references, metrics) for key, val in value.items()}


def format_author_block(paragraph):
    """Keep author, affiliation and correspondence in distinct native paragraphs."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    author = paragraph.insert_paragraph_before(style=paragraph.style)
    affiliation = paragraph.insert_paragraph_before(style=paragraph.style)
    author.add_run("Arjun Bajpai")
    author.add_run("1").font.superscript = True
    author.add_run("*").font.superscript = False
    affiliation.add_run("1").font.superscript = True
    affiliation.add_run("Windermere Preparatory School, Windermere, Florida, United States")
    paragraph.clear()
    paragraph.add_run("*Corresponding author: arjunbajpai2009@gmail.com")
    for item in (author, affiliation, paragraph):
        item.alignment = WD_ALIGN_PARAGRAPH.CENTER
        item.paragraph_format.line_spacing = 1
        item.paragraph_format.first_line_indent = Pt(0)
        item.paragraph_format.space_before = Pt(0)
        item.paragraph_format.space_after = Pt(0)
        if item is not paragraph:
            item.paragraph_format.keep_with_next = True
        for run in item.runs:
            run.font.name = "Times New Roman"


def format_docx(path, title, *, expected_reference_count):
    if type(expected_reference_count) is not int or expected_reference_count < 1:
        raise ValueError("Expected reference count must be a positive integer from the compiled bibliography")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from lxml import etree
    doc = Document(path)
    known_styles = {style.style_id for style in doc.styles}
    for style in doc.element.xpath(".//w:pStyle"):
        if style.get(qn("w:val")) not in known_styles:
            style.set(qn("w:val"), "Normal")
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.first_line_indent = Pt(0)
            style.font.size = Pt(12)
            if style.name.startswith("Heading"):
                style.font.bold = True
                style.paragraph_format.keep_with_next = True
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name in doc.styles:
            doc.styles[name].paragraph_format.space_before = Pt(0)
            doc.styles[name].paragraph_format.space_after = Pt(0)
    for style in doc.styles:
        pr = style.element.find(qn("w:pPr"))
        if pr is not None:
            for border in list(pr.findall(qn("w:pBdr"))):
                pr.remove(border)
    # Pandoc's table metadata and undefined Compact style are poorly handled by
    # some Word readers. Rebuild the native table grid with python-docx while
    # preserving each cell's paragraphs and runs.
    for old in list(doc.tables):
        column_layouts = {3: [1.1, 2.55, 2.85], 6: [.8, 1.05, .75, 1.8, 1.05, 1.05]}
        if len(old.columns) not in column_layouts:
            raise ValueError("Unreviewed table column count in Online export: " + str(len(old.columns)))
        new = doc.add_table(rows=len(old.rows), cols=len(old.columns))
        widths = column_layouts[len(old.columns)]
        if len(old.columns) == 6 and old.rows[0].cells[2].text.strip() == "Mean":
            widths = [.8, 1.05, .85, 1.35, 1.35, 1.1]
        for col, width in zip(new.columns, widths):
            col.width = Inches(width)
        for old_row, new_row in zip(old.rows, new.rows):
            for old_cell, new_cell, width in zip(old_row.cells, new_row.cells, widths):
                new_cell.width = Inches(width)
                for paragraph in list(new_cell._tc.findall(qn("w:p"))):
                    new_cell._tc.remove(paragraph)
                for paragraph in old_cell.paragraphs:
                    element = copy.deepcopy(paragraph._p)
                    for style in element.xpath("./w:pPr/w:pStyle"):
                        style.set(qn("w:val"), "Normal")
                    new_cell._tc.append(element)
        old._tbl.addprevious(new._tbl)
        old._tbl.getparent().remove(old._tbl)
    main_text = False
    previous_heading = False
    for paragraph in doc.paragraphs:
        heading = paragraph.style.name.startswith("Heading")
        if heading and paragraph.text != "Abstract":
            main_text = True
        caption = bool(re.match(r"(?:Figure|Table) \d+ \|", paragraph.text) or re.match(r"[AB]\) ", paragraph.text))
        equation = paragraph.text.lstrip().startswith(("$$", r"\begin{equation}", r"\begin{align}"))
        drawing = bool(paragraph._p.xpath(".//w:drawing"))
        numbered = bool(paragraph._p.xpath("./w:pPr/w:numPr")) or paragraph.style.name.startswith("List")
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.widow_control = True
        # First paragraphs, display material, and metadata are flush left.
        # Only continuing prose gets the 1.5em cue used by the LaTeX source.
        if not numbered:
            indent = main_text and not (heading or previous_heading or caption or equation or drawing)
            paragraph.paragraph_format.first_line_indent = Pt(18 if indent else 0)
        if paragraph.style.name in {"Title", "Subtitle", "Author", "Date"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif not heading:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2 if caption else 0)
        if "arjunbajpai2009@gmail.com" in paragraph.text:
            format_author_block(paragraph)
        if paragraph.text.startswith("Keywordsintegrated"):
            for run in paragraph.runs:
                if run.text == "Keywords":
                    run.text += " "
                    break
                if run.text.startswith("Keywordsintegrated"):
                    run.text = run.text.replace("Keywordsintegrated", "Keywords integrated", 1)
                    break
        if re.match(r"Table \d+ \|", paragraph.text) or re.match(r"[AB]\) ", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
        # Remove title borders inherited from templates without touching text.
        if paragraph.style.name == "Title":
            pr = paragraph._p.get_or_add_pPr()
            for border in list(pr.findall(qn("w:pBdr"))):
                pr.remove(border)
        if drawing:
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
        previous_heading = heading
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.lstrip().startswith(("$$", r"\begin{equation}", r"\begin{align}")):
            paragraph.paragraph_format.keep_together = True
            if index and paragraphs[index - 1].text.strip():
                paragraphs[index - 1].paragraph_format.keep_with_next = True
    for table in doc.tables:
        table.autofit = False
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            elem = OxmlElement("w:" + edge)
            elem.set(qn("w:val"), "single")
            elem.set(qn("w:sz"), "4")
            elem.set(qn("w:color"), "D9D9D9")
            borders.append(elem)
        old = props.find(qn("w:tblBorders"))
        if old is not None:
            props.remove(old)
        props.append(borders)
        for index, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Times New Roman", Pt(12)
                if index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EEEEEE")
                    cell._tc.get_or_add_tcPr().append(shade)
            if index == 0:
                header = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(header)
    doc.core_properties.title = title
    doc.core_properties.author = "Arjun Bajpai"
    doc.core_properties.subject = "Local Online citation manuscript"
    doc.core_properties.comments = "Equations retain TeX source under the NHSJS LaTeX export instructions."
    # Use explicit visible reference numbers. This avoids reliance on numbering
    # definitions inherited from a reference template's instructional examples.
    in_references, reference_number = False, 0
    for paragraph in doc.paragraphs:
        if paragraph.text == "References":
            in_references = True
            continue
        if in_references and paragraph.style.name.startswith("Heading"):
            in_references = False
        if in_references and paragraph.text.strip():
            reference_number += 1
            paragraph.paragraph_format.first_line_indent = Pt(0)
            pr = paragraph._p.get_or_add_pPr()
            for numbering in list(pr.findall(qn("w:numPr"))):
                pr.remove(numbering)
            run = paragraph.add_run(str(reference_number) + ". ")
            paragraph._p.insert(list(paragraph._p).index(pr) + 1, run._r)
    if reference_number != expected_reference_count:
        raise ValueError(
            f"Unexpected reference count in final Word document: "
            f"expected {expected_reference_count}, found {reference_number}"
        )
    # Pandoc copies image relationships from the reference file even when its
    # sample figures are absent from the new body. Keep only referenced images.
    referenced_ids = set(doc.element.xpath(".//@r:embed | .//@r:link | .//@r:id"))
    for relation_id, relation in list(doc.part.rels.items()):
        if relation.reltype.endswith("/image") and relation_id not in referenced_ids:
            doc.part.drop_rel(relation_id)
        elif relation.reltype.endswith("/comments") and not len(etree.fromstring(relation.target_part.blob)):
            doc.part.drop_rel(relation_id)
    doc.save(path)


def export(pandoc: Path, template: Path, intermediates: Path, output: Path):
    paper = ROOT / "paper"
    build = json.loads((intermediates / "build.json").read_text(encoding="utf-8"))
    if build["anonymous"] or build["manuscript_sha256"] != digest(paper / "paper.tex"):
        raise ValueError("Word export requires identified intermediates matching the current manuscript")
    for name in ("revision.aux", "revision.bbl"):
        if digest(intermediates / name) != build["intermediates_sha256"][name]:
            raise ValueError("Compiled intermediate changed " + name)
    source = select_identified(strip_pdf_layout_definitions(strip_comments((paper / "paper.tex").read_text(encoding="utf-8"))))
    source = normalize_table_stacks(expand_inputs(source, paper))
    source = source.replace(r"\textbf{Keywords}\quad", r"\textbf{Keywords}\ ")
    bbl = (intermediates / "revision.bbl").read_text(encoding="utf-8")
    entries = bibliography(bbl)
    references, plain_refs = {}, {}
    for key, value in entries.items():
        parsed, _ = pandoc_run(pandoc, value, "latex", "json")
        blocks = json.loads(parsed)["blocks"]
        inlines = []
        for block in blocks:
            if block["t"] not in {"Para", "Plain"}:
                raise ValueError("Unexpected compiled bibliography structure")
            inlines.extend(block["c"])
        references[key] = inlines
        plain_refs[key] = inline_text(inlines)
    aux = (intermediates / "revision.aux").read_text(encoding="utf-8")
    labels = dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^{}]+)\}", aux))
    def reference(match):
        if match.group(1) not in labels:
            raise ValueError("Unresolved cross-reference " + match.group(1))
        return labels[match.group(1)]
    source = re.sub(r"\\ref\{([^}]+)\}", reference, source)
    source = re.sub(r"\\eqref\{([^}]+)\}", lambda m: "(" + reference(m) + ")", source)
    ref_tex = "\\section*{References}\n\\begin{enumerate}\n" + "\n".join(r"\item " + value for value in entries.values()) + "\n\\end{enumerate}"
    source = re.sub(r"\\bibliographystyle\{[^}]+\}", "", source)
    source = re.sub(r"\\bibliography\{[^}]+\}", lambda _: ref_tex, source)
    source = re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\.pdf\}", lambda m: r"\includegraphics" + (m.group(1) or "") + "{" + m.group(2) + ".png}", source)
    ast_text, warnings = pandoc_run(pandoc, source, "latex", "json", "--resource-path", paper)
    ast = json.loads(ast_text)
    ast["blocks"] = prepare_blocks(ast["blocks"])
    number_headings(ast["blocks"], labels)
    seen_displays = set()
    for block in ast["blocks"]:
        if block["t"] in {"Figure", "Table"}:
            label = block["c"][0][0]
            caption = block["c"][1][1]
            if label in seen_displays:
                block["c"][1][1] = []
            elif caption:
                name = "Figure" if block["t"] == "Figure" else "Table"
                caption[0]["c"] = [{"t": "Strong", "c": text_inlines(name + " " + labels[label] + " | ")}] + caption[0]["c"]
            seen_displays.add(label)
    title = inline_text(ast.get("meta", {}).get("title", {}).get("c", []))
    metrics = {"citation_uses": [], "inline_equations": 0, "display_equations": 0, "tables": 0, "images": []}
    ast = transform_ast(ast, references, metrics)
    if not metrics["citation_uses"] or metrics["tables"] < 2:
        raise ValueError("Online export lost citations or native tables")
    output.parent.mkdir(parents=True, exist_ok=True)
    _, write_warnings = pandoc_run(pandoc, json.dumps(ast), "json", "docx", "--reference-doc", template,
                                   "--resource-path", paper, "--output", output)
    format_docx(output, title, expected_reference_count=len(entries))
    report = {"schema_version": 1, "source_sha256": digest(paper / "paper.tex"), "bibliography_sha256": digest(intermediates / "revision.bbl"),
              "reference_template_sha256": digest(template), "output_sha256": digest(output), "title": title,
              "representation": "Online citations and literal TeX equation source required by the journal's LaTeX route",
              "metrics": metrics, "references": plain_refs, "reader_warnings": warnings, "writer_warnings": write_warnings}
    output.with_suffix(".export.json").write_text(json.dumps(report, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    output.with_suffix(".content.json").write_text(json.dumps(ast, ensure_ascii=True) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "metrics": metrics, "warnings": warnings + write_warnings}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pandoc", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--intermediates", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    export(args.pandoc.resolve(), args.template.resolve(), args.intermediates.resolve(), args.output.resolve())
